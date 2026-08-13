"""
NSIP KM1 Reconciliation Script
Cross-checks cleaned NSIP KM1 Excel against existing database records.
READ-ONLY - does not modify database or Excel.
"""
import openpyxl
import mysql.connector
from collections import Counter, defaultdict

EXCEL_PATH = "NSIPKM1_12.08.2026.xlsx"
DB_CONFIG = {
    "host": "localhost",
    "port": 3306,
    "database": "nct_stocks_inventory",
    "user": "nctuser",
    "password": "Azrin@0019",
}

# ---------- 1. Read Excel ----------
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb["NSIP KM1"]

headers = []
excel_rows = []
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True), 1):
    if i == 1:
        headers = [str(h).strip() if h else "" for h in row]
    else:
        excel_rows.append(dict(zip(headers, row)))

print(f"=== EXCEL DATA ===")
print(f"Headers: {headers}")
print(f"Total Excel records: {len(excel_rows)}")

# Normalize Excel records
def norm(val):
    if val is None:
        return ""
    return str(val).strip()

def norm_upper(val):
    return norm(val).upper()

excel_units = []
for r in excel_rows:
    unit = {
        "unit_no": norm_upper(r.get("Unit No.", r.get("Unit No", ""))),
        "project": norm_upper(r.get("Project", "")),
        "phase_name": norm_upper(r.get("Phase Name", r.get("Phase", ""))),
        "block_name": norm_upper(r.get("Block Name", r.get("Block", ""))),
        "status": norm_upper(r.get("Status", "")),
        "unit_type": norm_upper(r.get("Unit Type", "")),
        "list_price": r.get("List Price", r.get("Listing Price", 0)),
        "built_up_area": r.get("Built-up Area", r.get("Built Up Area", r.get("Built_Up", 0))),
        "land_area": r.get("Land Area", r.get("Land_Area", 0)),
    }
    excel_units.append(unit)

print(f"\nExcel status counts: {dict(Counter(u['status'] for u in excel_units))}")
print(f"Excel project counts: {dict(Counter(u['project'] for u in excel_units))}")
print(f"Unique unit numbers in Excel: {len(set(u['unit_no'] for u in excel_units))}")

# Check for duplicates in Excel
excel_unit_counts = Counter(u['unit_no'] for u in excel_units)
excel_dups = {k: v for k, v in excel_unit_counts.items() if v > 1}
print(f"Excel duplicate unit numbers: {excel_dups if excel_dups else 'None'}")

# ---------- 2. Read Database ----------
conn = mysql.connector.connect(**DB_CONFIG)
cursor = conn.cursor(dictionary=True)

# Get all NSIP KM1 records
cursor.execute("""
    SELECT * FROM units_master 
    WHERE project LIKE %s
""", ("%NCT SMART INDUSTRIAL PARK%",))
db_rows = cursor.fetchall()
print(f"\n=== DATABASE DATA ===")
print(f"Total DB records matching 'NCT SMART INDUSTRIAL PARK': {len(db_rows)}")

db_units = []
for r in db_rows:
    unit = {
        "id": r.get("id"),
        "unit_no": norm_upper(r.get("unit_no", "")),
        "project": norm_upper(r.get("project", "")),
        "phase_name": norm_upper(r.get("phase_name", "")),
        "block_name": norm_upper(r.get("block_name", "")),
        "status": norm_upper(r.get("status", "")),
        "unit_type": norm_upper(r.get("unit_type", "")),
        "list_price": r.get("list_price", 0),
        "built_up_area": r.get("built_up_area", 0),
        "land_area": r.get("land_area", 0),
    }
    db_units.append(unit)

print(f"DB status counts: {dict(Counter(u['status'] for u in db_units))}")
print(f"DB project counts: {dict(Counter(u['project'] for u in db_units))}")
print(f"Unique unit numbers in DB: {len(set(u['unit_no'] for u in db_units))}")

# Check for duplicates in DB
db_unit_counts = Counter(u['unit_no'] for u in db_units)
db_dups = {k: v for k, v in db_unit_counts.items() if v > 1}
print(f"DB duplicate unit numbers: {db_dups if db_dups else 'None'}")

# ---------- 3. Build match keys ----------
def build_key(u):
    """Build match key from Project + Phase + Block + Unit No."""
    return (u["project"], u["phase_name"], u["block_name"], u["unit_no"])

excel_keys = {}
for u in excel_units:
    key = build_key(u)
    if key in excel_keys:
        excel_keys[key].append(u)
    else:
        excel_keys[key] = [u]

db_keys = {}
for u in db_units:
    key = build_key(u)
    if key in db_keys:
        db_keys[key].append(u)
    else:
        db_keys[key] = [u]

# ---------- 4. Comparison ----------
# A. Excel units not in DB
excel_not_in_db = []
for key, units in excel_keys.items():
    if key not in db_keys:
        for u in units:
            excel_not_in_db.append(u)

# B. DB units not in Excel
db_not_in_excel = []
for key, units in db_keys.items():
    if key not in excel_keys:
        for u in units:
            db_not_in_excel.append(u)

# C. Status changes (units in both)
status_changes = []
for key in set(excel_keys.keys()) & set(db_keys.keys()):
    excel_units_for_key = excel_keys[key]
    db_units_for_key = db_keys[key]
    # Match by unit_no within the key
    for eu in excel_units_for_key:
        for du in db_units_for_key:
            if eu["unit_no"] == du["unit_no"]:
                if eu["status"] != du["status"]:
                    status_changes.append({
                        "unit_no": eu["unit_no"],
                        "db_status": du["status"],
                        "excel_status": eu["status"],
                        "db_id": du["id"],
                    })

# D. Other field differences
field_diffs = []
for key in set(excel_keys.keys()) & set(db_keys.keys()):
    excel_units_for_key = excel_keys[key]
    db_units_for_key = db_keys[key]
    for eu in excel_units_for_key:
        for du in db_units_for_key:
            if eu["unit_no"] == du["unit_no"]:
                diffs = []
                # List Price
                try:
                    ep = float(eu["list_price"]) if eu["list_price"] else 0
                    dp = float(du["list_price"]) if du["list_price"] else 0
                    if abs(ep - dp) > 0.01:
                        diffs.append(("List Price", dp, ep))
                except (ValueError, TypeError):
                    pass
                # Unit Type
                if eu["unit_type"] != du["unit_type"]:
                    diffs.append(("Unit Type", du["unit_type"], eu["unit_type"]))
                # Built-up Area
                try:
                    eb = float(eu["built_up_area"]) if eu["built_up_area"] else 0
                    db_ = float(du["built_up_area"]) if du["built_up_area"] else 0
                    if abs(eb - db_) > 0.01:
                        diffs.append(("Built-up Area", db_, eb))
                except (ValueError, TypeError):
                    pass
                # Land Area
                try:
                    el = float(eu["land_area"]) if eu["land_area"] else 0
                    dl = float(du["land_area"]) if du["land_area"] else 0
                    if abs(el - dl) > 0.01:
                        diffs.append(("Land Area", dl, el))
                except (ValueError, TypeError):
                    pass
                # Phase Name
                if eu["phase_name"] != du["phase_name"]:
                    diffs.append(("Phase Name", du["phase_name"], eu["phase_name"]))
                # Block Name
                if eu["block_name"] != du["block_name"]:
                    diffs.append(("Block Name", du["block_name"], eu["block_name"]))
                # Project
                if eu["project"] != du["project"]:
                    diffs.append(("Project", du["project"], eu["project"]))
                
                if diffs:
                    field_diffs.append({
                        "unit_no": eu["unit_no"],
                        "db_id": du["id"],
                        "diffs": diffs,
                    })

# ---------- 5. Output Report ----------
print("\n" + "="*80)
print("NSIP KM1 RECONCILIATION REPORT")
print("="*80)

print(f"\n--- A. NEW EXCEL UNITS NOT IN DATABASE ({len(excel_not_in_db)}) ---")
if excel_not_in_db:
    for u in sorted(excel_not_in_db, key=lambda x: x["unit_no"]):
        print(f"  {u['unit_no']} | Status: {u['status']} | Project: {u['project']} | Phase: {u['phase_name']} | Block: {u['block_name']} | Type: {u['unit_type']} | Price: {u['list_price']}")
else:
    print("  None")

print(f"\n--- B. DATABASE UNITS MISSING FROM EXCEL ({len(db_not_in_excel)}) ---")
if db_not_in_excel:
    for u in sorted(db_not_in_excel, key=lambda x: x["unit_no"]):
        print(f"  {u['unit_no']} | Status: {u['status']} | Project: {u['project']} | Phase: {u['phase_name']} | Block: {u['block_name']} | Type: {u['unit_type']} | Price: {u['list_price']} | DB ID: {u['id']}")
else:
    print("  None")

print(f"\n--- C. STATUS CHANGES ({len(status_changes)}) ---")
if status_changes:
    for sc in sorted(status_changes, key=lambda x: x["unit_no"]):
        direction = ""
        if sc["db_status"] in ("SIGNED", "SOLD", "REGISTERED") and sc["excel_status"] == "AVAILABLE":
            direction = "  <-- Previously Signed/Sold -> now Available"
        elif sc["db_status"] == "AVAILABLE" and sc["excel_status"] in ("SIGNED", "SOLD", "REGISTERED"):
            direction = "  <-- Previously Available -> now Signed/Sold"
        print(f"  {sc['unit_no']} | DB: {sc['db_status']} -> Excel: {sc['excel_status']}{direction}")
else:
    print("  None")

print(f"\n--- D. OTHER FIELD DIFFERENCES ({len(field_diffs)}) ---")
if field_diffs:
    for fd in sorted(field_diffs, key=lambda x: x["unit_no"]):
        print(f"  Unit: {fd['unit_no']} (DB ID: {fd['db_id']})")
        for field, old_val, new_val in fd["diffs"]:
            print(f"    {field}: DB={old_val} -> Excel={new_val}")
else:
    print("  None")

# ---------- 6. Summary ----------
print(f"\n--- E. RECONCILIATION SUMMARY ---")
print(f"  Existing NSIP KM1 database records: {len(db_units)}")
print(f"  New Excel records: {len(excel_units)}")
print(f"  Exact matches (same key in both): {len(set(excel_keys.keys()) & set(db_keys.keys()))}")
print(f"  Missing/new records (Excel not in DB): {len(excel_not_in_db)}")
print(f"  DB records missing from Excel: {len(db_not_in_excel)}")
print(f"  Status changes: {len(status_changes)}")
print(f"  Other field differences: {len(field_diffs)}")
print(f"  Excel duplicate unit numbers: {len(excel_dups)}")
print(f"  DB duplicate unit numbers: {len(db_dups)}")

# Check expected units
print(f"\n--- EXPECTED UNITS CHECK ---")
expected = {"C35", "D93", "D97", "H26"}
for unit_no in sorted(expected):
    # Find in Excel
    excel_match = [u for u in excel_units if u["unit_no"] == unit_no]
    db_match = [u for u in db_units if u["unit_no"] == unit_no]
    print(f"  {unit_no}:")
    if excel_match:
        print(f"    Excel: Status={excel_match[0]['status']}, Phase={excel_match[0]['phase_name']}, Block={excel_match[0]['block_name']}")
    else:
        print(f"    Excel: NOT FOUND")
    if db_match:
        print(f"    DB: Status={db_match[0]['status']}, Phase={db_match[0]['phase_name']}, Block={db_match[0]['block_name']}, ID={db_match[0]['id']}")
    else:
        print(f"    DB: NOT FOUND")

cursor.close()
conn.close()
print("\n=== RECONCILIATION COMPLETE ===")